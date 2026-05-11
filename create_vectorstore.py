import os
import openpyxl
from dotenv import load_dotenv
from langchain_core.documents import Document
from langchain_community.vectorstores import FAISS

# Import embeddings from the rag_agent if possible, or initialize here
from langchain_google_genai import GoogleGenerativeAIEmbeddings

load_dotenv()

# ── Constants ─────────────────────────────────────────────────────────────────
EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PACT_DUMMY_DATA (1) 1.xlsx")
EXCEL_SHEET = 0          # sheet index (0 = first sheet) or sheet name string
FAISS_INDEX_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "faiss_index")

def _row_to_text(sheet_name: str, columns: list[str], row: tuple) -> str:
    """Convert an Excel row into a human-readable string for embedding."""
    parts = [f"Sheet: {sheet_name}"] + [
        f"{col}: {val}" for col, val in zip(columns, row)
    ]
    return " | ".join(parts)

def create_and_save_vectorstore():
    print(f"⏳  Building RAG vectorstore from Excel: {EXCEL_PATH} …")

    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    # Resolve sheet by index or name
    if isinstance(EXCEL_SHEET, int):
        ws = wb.worksheets[EXCEL_SHEET]
    else:
        ws = wb[EXCEL_SHEET]

    sheet_name = ws.title
    rows_iter = ws.iter_rows(values_only=True)

    # First row = column headers
    raw_headers = next(rows_iter)
    columns = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(raw_headers)]

    docs = []
    for raw_row in rows_iter:
        row = tuple("" if v is None else str(v) for v in raw_row)
        content = _row_to_text(sheet_name, columns, row)
        docs.append(Document(page_content=content, metadata={"sheet": sheet_name}))

    wb.close()

    print(f"✅  Created {len(docs)} documents from sheet '{sheet_name}'.")
    print(f"Let's embed and save to FAISS index...")

    embeddings = GoogleGenerativeAIEmbeddings(
        model="models/gemini-embedding-001",
        google_api_key=os.getenv("GOOGLE_API_KEY"),
    )
    
    vectorstore = FAISS.from_documents(docs, embeddings)
    vectorstore.save_local(FAISS_INDEX_PATH)
    print(f"🎉 Successfully saved FAISS index to {FAISS_INDEX_PATH}")

if __name__ == "__main__":
    create_and_save_vectorstore()
